"""
Excel Data Reader for Card Printing
Reads card holder information from Excel files using openpyxl
"""

import openpyxl
import json

def load_config():
    """Load configuration from config.json"""
    with open('config.json', 'r') as f:
        return json.load(f)

def read_card_data(excel_file=None, sheet_name=None):
    """
    Read card data from Excel file
    
    Expected Excel format:
    | Name | ID | Department | Photo | Title | Email |
    |------|-----|-----------|-------|-------|-------|
    | John | 001 | Sales     | photo1.jpg | Manager | john@example.com |
    
    Returns: List of dictionaries with card data
    """
    config = load_config()
    
    if excel_file is None:
        excel_file = config['excel']['data_file']
    if sheet_name is None:
        sheet_name = config['excel']['sheet_name']
    
    print(f"Reading card data from: {excel_file}, Sheet: {sheet_name}")
    
    try:
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook[sheet_name]
        
        # Get headers from first row
        headers = []
        for cell in sheet[1]:
            if cell.value:
                headers.append(cell.value)
        
        print(f"Found columns: {headers}")
        
        # Read data rows
        card_data = []
        for row_idx in range(2, sheet.max_row + 1):
            row_data = {}
            for col_idx, header in enumerate(headers, start=1):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                row_data[header] = cell_value if cell_value is not None else ""
            
            # Only add rows that have at least a name
            if row_data.get('Name'):
                card_data.append(row_data)
        
        print(f"Loaded {len(card_data)} card records")
        return card_data
        
    except FileNotFoundError:
        print(f"Error: Excel file '{excel_file}' not found!")
        print("Please create an Excel file with card data.")
        return []
    except KeyError:
        print(f"Error: Sheet '{sheet_name}' not found in Excel file!")
        print(f"Available sheets: {workbook.sheetnames}")
        return []
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return []

def create_sample_excel():
    """Create a sample Excel file with example card data"""
    config = load_config()
    excel_file = config['excel']['data_file']
    sheet_name = config['excel']['sheet_name']
    
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    
    # Add headers
    headers = ['Name', 'ID', 'Department', 'Photo', 'Title', 'Email']
    sheet.append(headers)
    
    # Add sample data
    sample_data = [
        ['John Doe', '001', 'Sales', 'photos/john.jpg', 'Sales Manager', 'john@company.com'],
        ['Jane Smith', '002', 'IT', 'photos/jane.jpg', 'Developer', 'jane@company.com'],
        ['Bob Johnson', '003', 'HR', 'photos/bob.jpg', 'HR Director', 'bob@company.com'],
    ]
    
    for row in sample_data:
        sheet.append(row)
    
    # Auto-adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width
    
    workbook.save(excel_file)
    print(f"Sample Excel file created: {excel_file}")
    print("Edit this file with your actual card data.")


if __name__ == "__main__":
    print("=" * 60)
    print("Excel Card Data Reader")
    print("=" * 60)
    
    # Create sample Excel if it doesn't exist
    import os
    config = load_config()
    if not os.path.exists(config['excel']['data_file']):
        print("Excel file not found. Creating sample...")
        create_sample_excel()
    
    # Read and display card data
    cards = read_card_data()
    
    print("\n" + "=" * 60)
    print("Card Data Preview:")
    print("=" * 60)
    
    for idx, card in enumerate(cards, 1):
        print(f"\nCard {idx}:")
        for key, value in card.items():
            print(f"  {key}: {value}")
    
    print("\n" + "=" * 60)